"""Jira import Excel: Global Search Bar → SPACE application integration (to 31-Aug-2026)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = ["Key", "Type", "Component", "Title", "Est. Hours", "Description"]

# Key left empty (optional). Type = Story. Component = global search bar.
# Hours sized for ~2 remaining weeks (20–31 Aug) plus a compact "already done" story.
TICKETS = [
    {
        "title": "Document completed Global Search POC baseline for SPACE integration",
        "hours": 4,
        "description": (
            "Capture what is already live in the Global Search POC so SPACE integration "
            "does not re-do that work.\n\n"
            "Already done:\n"
            "- Multi-source search (BuildTrack, Assets, Space Action + sibling Space DBs)\n"
            "- Azure AI Search as discovery layer; SQL as source of truth; Top-K + LLM answers\n"
            "- Space Action schema-aware indexing (action.* tables, text columns only; "
            "excluded OfflineSyncLog / Synonyms_Configuration; ~6.6k docs in space-action-index)\n"
            "- Production flags: hydrate from SQL, catalog, incremental sync, security trim, metrics\n"
            "- Space ProjectId cross-DB hops (Project, HSE, Incident, LMS, CMS, Dashboard)\n"
            "- SearchBar types + Linked badge for related Space records\n"
            "- APIs: POST /api/search/results, POST /api/search/answer, GET /api/search/sources\n"
            "- Widget client under frontend/lib/global-search (createSearchClient)\n\n"
            "Deliverable: 1-page baseline note for SPACE app team (architecture, sources, APIs)."
        ),
    },
    {
        "title": "Agree SPACE embed strategy: widget vs iframe vs shared package",
        "hours": 8,
        "description": (
            "Workshop with SPACE application team to pick how Global Search Bar is hosted.\n\n"
            "Decide:\n"
            "- Embed SearchBar in SPACE header (shared React) vs iframe vs copy of lib/global-search\n"
            "- Auth: reuse SPACE session / SSO vs pass JWT to Global Search API\n"
            "- API base URL, CORS, DEV vs UAT vs PROD\n"
            "- Which Space SQL databases must be on for go-live (Action, Project, HSE, etc.)\n\n"
            "Deliverable: short integration design (embed choice, auth, env matrix)."
        ),
    },
    {
        "title": "Freeze Global Search API contract for SPACE host app",
        "hours": 6,
        "description": (
            "Publish a stable contract SPACE can call without depending on BuildTrack UI.\n\n"
            "Include:\n"
            "- POST /api/search/results (fast hits, no LLM)\n"
            "- POST /api/search/answer (Top-K + AI summary)\n"
            "- GET /api/search/sources\n"
            "- Request fields: query, filter_type, top_k, source_ids\n"
            "- Response fields: results[], metadata.database_id, content_type, url, score\n"
            "- Auth header (Bearer) and 401/422/503 behaviour\n\n"
            "Deliverable: API contract page + example payloads SPACE can implement against."
        ),
    },
    {
        "title": "Package SearchBar as reusable module with SPACE-ready props",
        "hours": 8,
        "description": (
            "Extract frontend/lib/global-search so SPACE can mount it without forking BuildTrack.\n\n"
            "Props/config:\n"
            "- apiBaseUrl, getAuthToken, theme tokens (SPACE colours)\n"
            "- contentTypes map (project, event, HSE, incident, document, etc.)\n"
            "- onOpenResult callback for SPACE routing\n"
            "- source_ids default for SPACE-only search\n\n"
            "Deliverable: reusable SearchBar module + README for SPACE developers."
        ),
    },
    {
        "title": "Map SPACE session user into search security and project scope",
        "hours": 8,
        "description": (
            "Wire SPACE logged-in user into Global Search so results respect SPACE permissions.\n\n"
            "Work:\n"
            "- Pass SPACE JWT/session token to GET/POST /api/search/*\n"
            "- Map user id / roles to search security_trim (project membership)\n"
            "- Pass active SPACE project/company as filter so queries are scoped\n"
            "- Fail closed on missing token; document empty vs forbidden states\n\n"
            "Deliverable: auth mapping table + working DEV login from SPACE shell."
        ),
    },
    {
        "title": "Mount Global Search Bar in SPACE application shell (DEV)",
        "hours": 8,
        "description": (
            "Embed Global Search Bar in SPACE header/shell on DEV.\n\n"
            "Work:\n"
            "- Place bar in SPACE chrome (header or command palette)\n"
            "- Match SPACE visual tokens where required\n"
            "- Point to Global Search API DEV (Space SQL via VPN / on-prem)\n"
            "- Smoke test: inspection / EVT number / list active projects / Linked hops\n\n"
            "Deliverable: SearchBar visible and usable inside SPACE DEV."
        ),
    },
    {
        "title": "Deep-link search results into SPACE module screens",
        "hours": 8,
        "description": (
            "Clicking a search result must open the correct SPACE screen, not BuildTrack routes.\n\n"
            "Map hit types to SPACE URLs, including:\n"
            "- space-project / project → SPACE project page\n"
            "- event / event_agenda / event_task → SPACE Action/event screens\n"
            "- HSE occupational/environmental review → SPACE HSE\n"
            "- incident → SPACE Incident\n"
            "- employment / LMS → SPACE LMS\n"
            "- contract / CMS → SPACE CMS\n"
            "- documents / DMS / attachment → SPACE DMS where keys exist\n\n"
            "Deliverable: url mapper + verified clicks for top content types."
        ),
    },
    {
        "title": "Harden Space ProjectId / DocumentId hops for SPACE demo queries",
        "hours": 6,
        "description": (
            "POC already hops ProjectId across Project, HSE, Incident, LMS, CMS, Dashboard. "
            "For SPACE integration, lock demo queries and fill remaining high-value hops.\n\n"
            "Work:\n"
            "- Regression pack: list active projects, HSE reviews, incidents, contracts, EVT/agenda\n"
            "- Add DocumentId hop (DMS/Attachment ↔ project) where FKs exist\n"
            "- Reduce duplicate schema_sql vs space_project cards in UI\n"
            "- Document which of 17 Space DBs are searchable vs hop-only\n\n"
            "Deliverable: demo query list + expected Linked badges."
        ),
    },
    {
        "title": "Deploy Global Search API for SPACE DEV with runbook",
        "hours": 8,
        "description": (
            "Stand up Global Search API for SPACE DEV (and note UAT path).\n\n"
            "Work:\n"
            "- Env: Space SQL URLs, Azure AI Search indexes, Azure OpenAI, CORS for SPACE origin\n"
            "- Schema catalog / space-index sync runbook (POST /api/search/space-index)\n"
            "- Health: /api/health, /api/search/space-index/status, /api/search/sources\n"
            "- Secrets: Key Vault or equivalent; no keys in SPACE frontend\n"
            "- One-process-per-port notes (avoid duplicate uvicorn on 8000)\n\n"
            "Deliverable: DEV deploy + 1-page runbook for SPACE ops."
        ),
    },
    {
        "title": "SPACE UAT checklist, latency baseline, and go/no-go for 31 Aug",
        "hours": 8,
        "description": (
            "UAT and readiness checkpoint before 31-Aug-2026.\n\n"
            "Work:\n"
            "- Checklist: login from SPACE, results, AI answer, Linked hops, empty query, 401\n"
            "- Measure /search/results vs /search/answer latency on SPACE DEV\n"
            "- Fix P1 UAT issues (auth, CORS, wrong deep-links, missing Space source)\n"
            "- Go/no-go: CORS, VPN/SQL, flags, rollback, demo script for manager\n"
            "- Backlog after 31 Aug: UserId hops, BMS name match, UAT/PROD full rollout\n\n"
            "Deliverable: UAT report + go/no-go note dated on or before 31 Aug."
        ),
    },
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Jira Stories"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, t in enumerate(TICKETS, 2):
        ws.cell(i, 1, "").border = thin  # Key optional
        type_cell = ws.cell(i, 2, "Story")
        type_cell.border = thin
        type_cell.alignment = center
        comp = ws.cell(i, 3, "global search bar")
        comp.border = thin
        comp.alignment = wrap
        title = ws.cell(i, 4, t["title"])
        title.border = thin
        title.alignment = wrap
        hours = ws.cell(i, 5, t["hours"])
        hours.border = thin
        hours.alignment = center
        desc = ws.cell(i, 6, t["description"])
        desc.border = thin
        desc.alignment = wrap
        ws.row_dimensions[i].height = 95

    widths = [12, 12, 22, 62, 12, 78]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(TICKETS) + 1}"

    # Plan sheet
    ws2 = wb.create_sheet("Integration Plan to 31 Aug")
    ws2["A1"] = "Global Search Bar → SPACE application — plan through 31 Aug 2026"
    ws2["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws2.merge_cells("A1:D1")

    plan_headers = ["When", "Focus", "Stories (titles)", "Hours"]
    for c, h in enumerate(plan_headers, 1):
        cell = ws2.cell(3, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin

    plan_rows = [
        (
            "Already done (before this sprint)",
            "POC search + Space data + ProjectId hops",
            "Document completed Global Search POC baseline for SPACE integration",
            4,
        ),
        (
            "20–22 Aug",
            "Design + API + widget package",
            "Agree embed strategy; Freeze API contract; Package SearchBar module",
            22,
        ),
        (
            "24–26 Aug",
            "Auth + embed in SPACE DEV",
            "Map SPACE session into search; Mount SearchBar in SPACE shell (DEV)",
            16,
        ),
        (
            "27–28 Aug",
            "Deep-links + demo hops",
            "Deep-link into SPACE modules; Harden ProjectId/DocumentId hops",
            14,
        ),
        (
            "29–31 Aug",
            "Deploy + UAT + go/no-go",
            "Deploy API for SPACE DEV; UAT checklist, latency, go/no-go by 31 Aug",
            16,
        ),
    ]
    for r_i, row in enumerate(plan_rows, 4):
        for c, val in enumerate(row, 1):
            cell = ws2.cell(r_i, c, val)
            cell.border = thin
            cell.alignment = wrap
        ws2.row_dimensions[r_i].height = 48

    ws2.column_dimensions["A"].width = 36
    ws2.column_dimensions["B"].width = 38
    ws2.column_dimensions["C"].width = 72
    ws2.column_dimensions["D"].width = 10

    total = sum(t["hours"] for t in TICKETS)
    ws2["A10"] = f"Total estimated hours (10 stories): {total}"
    ws2["A10"].font = Font(bold=True, color="1F4E79")
    ws2["A11"] = (
        "Type = Story for all rows. Component = global search bar. Key left blank (optional)."
    )
    ws2["A12"] = (
        "Out of scope until after 31 Aug: full UAT/PROD rollout, UserId graph hops, "
        "BMS name-match, re-enabling BuildTrack/Assets inside SPACE if not required for MVP."
    )
    ws2["A12"].alignment = wrap
    ws2.merge_cells("A12:C12")
    ws2.row_dimensions[12].height = 36

    out = "jira-space-integration-stories.xlsx"
    wb.save(out)
    print(f"Wrote {out}")
    print(f"Stories: {len(TICKETS)} | Hours: {total}")


if __name__ == "__main__":
    main()
