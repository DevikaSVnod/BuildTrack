"""Generate Global Search Bar worklog + SPACE client integration plan Excel."""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADERS = [
    "Ticket ID",
    "Component",
    "Log Date",
    "Day",
    "Hours",
    "Phase",
    "Title",
    "Task Description",
    "Status",
]

# (ticket, date, hours, phase, title, description, status)
ROWS = [
    (
        "GS-SPACE-01",
        date(2026, 8, 17),
        2.0,
        "Space Cross-DB Linking",
        "ProjectId hub design & key promotion",
        "Defined Space cross-DB linking hub on numeric ProjectId. Promoted "
        "ProjectId/UserId/DocumentId (and project name) from schema-SQL rows "
        "onto hit metadata in executor so related DBs can join on the same key.",
        "Done",
    ),
    (
        "GS-SPACE-02",
        date(2026, 8, 17),
        2.5,
        "Space Cross-DB Linking",
        "Cross-DB ProjectId hop fetchers (6 Space DBs)",
        "Implemented ProjectId hop fetchers in cross_db_links for Space Project, "
        "HSE (occupational + environmental reviews), Incident, LMS Employment, "
        "CMS ProjectCM, and Dashboard_Project_Details. Wired attach into multi-DB "
        "retrieve (including single Space source when siblings enabled).",
        "Done",
    ),
    (
        "GS-SPACE-03",
        date(2026, 8, 17),
        1.5,
        "Space Cross-DB Linking",
        "Display filter + SearchBar for linked Space types",
        "Fixed UI dropping linked hops: merge space_project_id related hits after "
        "answer/entity filters; added SearchBar TYPE_CONFIG for space_project, "
        "occupational/environmental review, incident, employment, contract, "
        "dashboard_project; stopped collapsing Space linked groups; Linked badge "
        "for related_via=space_project_id.",
        "Done",
    ),
    (
        "GS-SPACE-04",
        date(2026, 8, 17),
        1.5,
        "Space Cross-DB Linking",
        "SQL join-key injection (ProjectId in SELECT)",
        "Root-cause fix: LLM SQL for 'list active projects' omitted ProjectId. "
        "Validator now injects primary/join keys (ProjectId/UserId/DocumentId); "
        "generator prompt updated; related-hit budget and result slice increased. "
        "Verified ~34 linked results across types in UI.",
        "Done",
    ),
    (
        "GS-SPACE-05",
        date(2026, 8, 17),
        0.5,
        "Space Cross-DB Linking",
        "E2E verification & manager walkthrough prep",
        "Verified live /search/results and /search/answer for list active projects "
        "(ProjectIds + LINKED HSE/Incident/LMS/CMS). Documented which of 17 Space "
        "DBs are searchable vs ProjectId-linked for planning.",
        "Done",
    ),
    (
        "GS-SPACE-06",
        date(2026, 8, 18),
        4.0,
        "POC Hardening",
        "Expand ProjectId hops + regression test pack",
        "Add more ProjectId-linked tables where FKs exist; write a small regression "
        "pack (project list, HSE reviews, incidents, contracts) and document "
        "expected Linked badges / counts for demos.",
        "Planned",
    ),
    (
        "GS-SPACE-07",
        date(2026, 8, 19),
        4.0,
        "POC Hardening",
        "DMS / DocumentId linking (deferred hop)",
        "Design and implement DocumentId-based linking between Space DMS/Attachment "
        "and related project context so document queries can show Linked project/HSE "
        "where keys exist.",
        "Planned",
    ),
    (
        "GS-SPACE-08",
        date(2026, 8, 20),
        4.0,
        "POC Hardening",
        "UserId graph linking (deferred hop)",
        "Promote and hop on UserId across UM / LMS / Action where safe; keep "
        "permission boundaries clear for client readiness.",
        "Planned",
    ),
    (
        "GS-SPACE-09",
        date(2026, 8, 21),
        4.0,
        "POC Hardening",
        "Answer quality & Linked UX polish",
        "Tune AI answer citations for related_via=space_project_id; improve SearchBar "
        "grouping/tabs for Space content types; reduce duplicate schema_sql vs "
        "space_project cards where possible.",
        "Planned",
    ),
    (
        "GS-SPACE-10",
        date(2026, 8, 22),
        4.0,
        "POC Hardening",
        "Security trim + source toggles review",
        "Review security trimming for Space SQL results; confirm SEARCH_ENABLE_* "
        "flags and disable unused sources; document which DBs must be on for client demo.",
        "Planned",
    ),
    (
        "GS-SPACE-11",
        date(2026, 8, 25),
        8.0,
        "SPACE Client Integration",
        "Integration design with SPACE app team",
        "Align with SPACE (client) app architecture: embed vs iframe vs shared React "
        "package; auth (JWT/SSO/session cookie); API base URL; CORS; environment "
        "matrix (DEV/UAT/PROD). Produce short integration design note for manager.",
        "Planned",
    ),
    (
        "GS-SPACE-12",
        date(2026, 8, 26),
        8.0,
        "SPACE Client Integration",
        "Extract reusable Global Search Bar package/API contract",
        "Package SearchBar + API client as a reusable module (props: apiUrl, auth "
        "token, theme hooks). Freeze request/response contract for /search/results "
        "and /search/answer used by SPACE host app.",
        "Planned",
    ),
    (
        "GS-SPACE-13",
        date(2026, 8, 27),
        8.0,
        "SPACE Client Integration",
        "Auth & tenant/project context wiring",
        "Map SPACE logged-in user into search security context; pass active "
        "project/company filters into search so results respect SPACE session scope.",
        "Planned",
    ),
    (
        "GS-SPACE-14",
        date(2026, 8, 28),
        8.0,
        "SPACE Client Integration",
        "Embed Global Search Bar into SPACE shell (DEV)",
        "Mount Global Search Bar in SPACE application header/shell on DEV; match "
        "SPACE visual tokens where required; smoke-test list/filter queries against "
        "on-prem Space SQL.",
        "Planned",
    ),
    (
        "GS-SPACE-15",
        date(2026, 8, 29),
        6.0,
        "SPACE Client Integration",
        "Deep-link results into SPACE modules",
        "Map search hit types (project, HSE review, incident, employment, contract, "
        "document) to SPACE route URLs so clicking a result opens the correct SPACE screen.",
        "Planned",
    ),
    (
        "GS-SPACE-16",
        date(2026, 9, 1),
        8.0,
        "SPACE Client Integration",
        "Backend deploy for SPACE environments",
        "Deploy Global Search API (schema catalog sync, Azure OpenAI/AI Search config, "
        "Space SQL connectivity) for SPACE DEV/UAT; runbook for schema sync and health checks.",
        "Planned",
    ),
    (
        "GS-SPACE-17",
        date(2026, 9, 2),
        6.0,
        "SPACE Client Integration",
        "UAT test plan + performance baseline",
        "Write UAT checklist (queries per module, Linked hops, auth failures, empty "
        "states). Measure latency for /results vs /answer; set target SLAs for demo.",
        "Planned",
    ),
    (
        "GS-SPACE-18",
        date(2026, 9, 3),
        8.0,
        "SPACE Client Integration",
        "UAT fixes & Linked hop coverage gaps",
        "Fix UAT findings; add remaining high-value ProjectId (or DocumentId/UserId) "
        "hops requested by SPACE product owners.",
        "Planned",
    ),
    (
        "GS-SPACE-19",
        date(2026, 9, 4),
        4.0,
        "SPACE Client Integration",
        "Manager demo pack & handover",
        "Prepare demo script (active projects → Linked HSE/Incident/LMS), architecture "
        "one-pager (17 Space DBs + ProjectId hub), and backlog for post-go-live "
        "(BMS name match, Action/DMS linking, BuildTrack re-enable if needed).",
        "Planned",
    ),
    (
        "GS-SPACE-20",
        date(2026, 9, 5),
        8.0,
        "SPACE Client Integration",
        "SPACE UAT/PROD readiness checkpoint",
        "Final config review (env flags, secrets, CORS, VPN/SQL access), rollback "
        "plan, and go/no-go checklist with manager before wider SPACE rollout.",
        "Planned",
    ),
]

DAY_NAMES = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Global Search Worklog"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    done_fill = PatternFill("solid", fgColor="C6EFCE")
    planned_fill = PatternFill("solid", fgColor="FFEB9C")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_i, (ticket, d, hours, phase, title, desc, status) in enumerate(ROWS, 2):
        values = [
            ticket,
            "global search bar",
            d,
            DAY_NAMES[d.weekday()],
            hours,
            phase,
            title,
            desc,
            status,
        ]
        for c_i, val in enumerate(values, 1):
            cell = ws.cell(r_i, c_i, val)
            cell.border = thin
            cell.alignment = wrap
            if c_i == 3:
                cell.number_format = "DD-MMM-YYYY"
            if c_i == 9:
                cell.fill = done_fill if status == "Done" else planned_fill
                cell.alignment = Alignment(horizontal="center", vertical="top")

    widths = {
        "A": 14,
        "B": 18,
        "C": 13,
        "D": 6,
        "E": 8,
        "F": 24,
        "G": 42,
        "H": 72,
        "I": 10,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 22
    for r in range(2, len(ROWS) + 2):
        ws.row_dimensions[r].height = 70
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(ROWS) + 1}"

    # Summary sheet
    ws2 = wb.create_sheet("Plan Summary")
    ws2["A1"] = "Global Search Bar → SPACE Client — Plan Summary"
    ws2["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws2.merge_cells("A1:D1")

    summary_headers = ["Section", "Item", "Hours", "Notes"]
    for c, h in enumerate(summary_headers, 1):
        cell = ws2.cell(3, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin

    done_h = sum(r[2] for r in ROWS if r[6] == "Done")
    planned_h = sum(r[2] for r in ROWS if r[6] == "Planned")
    hardening = sum(r[2] for r in ROWS if r[3] == "POC Hardening")
    client = sum(r[2] for r in ROWS if r[3] == "SPACE Client Integration")

    summary_rows = [
        (
            "Completed (17-Aug-2026)",
            "Space ProjectId cross-DB linking + UI/SQL fixes",
            done_h,
            "Verified in BuildTrack POC UI: ~34 linked results",
        ),
        (
            "Near-term",
            "POC hardening (hops, DocumentId/UserId, UX, security)",
            hardening,
            "18–22 Aug (indicative — adjust with manager)",
        ),
        (
            "Future",
            "Integrate Global Search Bar into SPACE client app",
            client,
            "Design → embed → deep-link → UAT → readiness",
        ),
        (
            "Total (this sheet)",
            "All tickets GS-SPACE-01 … 20",
            done_h + planned_h,
            "Dates are planning estimates",
        ),
        (
            "Data connected today",
            "17 Space DBs searchable; 6 in ProjectId hub",
            "-",
            "Project, HSE, Incident, LMS, CMS, Dashboard",
        ),
        (
            "Later backlog",
            "BMS name match, Action/DMS deep link, BuildTrack/Assets re-enable",
            "-",
            "After SPACE embed MVP",
        ),
    ]

    for i, row in enumerate(summary_rows, 4):
        for c, val in enumerate(row, 1):
            cell = ws2.cell(i, c, val)
            cell.border = thin
            cell.alignment = wrap

    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 58
    ws2.column_dimensions["C"].width = 10
    ws2.column_dimensions["D"].width = 52
    for r in range(4, 10):
        ws2.row_dimensions[r].height = 35

    ws2["A11"] = 'Component for all task rows: "global search bar"'
    ws2["A11"].font = Font(italic=True, color="666666")
    ws2["A12"] = (
        "Prepared for manager planning — integrate Global Search Bar into "
        "SPACE application (client)."
    )
    ws2["A12"].font = Font(italic=True, color="666666")

    out = "global-search-bar-space-worklog-plan.xlsx"
    wb.save(out)
    print(f"Wrote {out}")
    print(f"Done hours: {done_h} | Planned hours: {planned_h} | Total: {done_h + planned_h}")


if __name__ == "__main__":
    main()
