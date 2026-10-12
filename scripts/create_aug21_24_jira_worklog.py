"""Jira worklog Excel: 3 stories, 21-24 Aug 2026 weekdays, 8h/day, component AI/ML."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = ["Key", "Type", "Component", "Title", "Est. Hours", "Description"]

# 2 weekdays x 8h = 16h across 3 stories (22-23 Aug are Sat/Sun)
TICKETS = [
    {
        "hours": 4,
        "title": "Review the Global Search POC and confirm what can be reused for SPACE",
        "description": (
            "Date: 21-Aug-2026 (Fri).\n\n"
            "Went through the full POC and traced one real query end to end, so we know exactly "
            "what we are carrying into SPACE and what we are leaving behind.\n\n"
            "What the POC is today:\n"
            "- FastAPI backend, Next.js frontend, Azure AI Search, Azure OpenAI, Azure Blob.\n"
            "- It opens 19 direct SQL connections (BuildTrack, Assets and 17 SPACE databases).\n"
            "- No LangChain or LangGraph is used, so the new work must not be planned around one.\n\n"
            "Reuse (this is the real value in the POC):\n"
            "- Query understanding and query planner - rule based, fast, no LLM call.\n"
            "- Ranking logic - the hand-tuned boosts are the actual search intelligence.\n"
            "- Hybrid search - keyword plus meaning plus vector, run together.\n"
            "- Azure index design, the catalog pattern, and the incremental sync (only re-embeds "
            "rows that actually changed, so it controls cost).\n\n"
            "Do not reuse, with the reason:\n"
            "- The POC's own login - it would create a second login next to SPACE.\n"
            "- The 19 direct database connections - they ignore who is searching.\n"
            "- The LLM-writes-SQL feature - it runs generated SQL on live data, on the typing path.\n"
            "- In-memory caches - not per user and do not work with more than one server.\n"
            "- Backend-generated page links - the frontend should own its own routes.\n\n"
            "Bottom line: about 25-30 percent of the POC is worth carrying forward."
        ),
    },
    {
        "hours": 4,
        "title": "Check user permissions and security in the POC search path",
        "description": (
            "Date: 21-Aug-2026 (Fri).\n\n"
            "Checked whether the POC can safely show different results to different users. This "
            "decides whether we can build on it or have to design fresh.\n\n"
            "Main findings:\n"
            "- Only two permission checks exist in the whole project, and the real one is switched "
            "off in the current setup. So no permission filtering runs today.\n"
            "- Two users with different access get exactly the same results.\n"
            "- The result cache does not include the user, so one user's results can be shown to "
            "another user for up to 2 minutes.\n"
            "- One place does check permission but then ignores the answer and keeps the row.\n"
            "- The AI answer is written from unfiltered results, so it can describe records the "
            "user is not allowed to see.\n"
            "- Reindex and clear-cache endpoints can be called by any logged-in user.\n"
            "- Live database passwords and Azure keys are sitting in backend/.env.\n\n"
            "Why this matters for the integration: this cannot be added on later as a wrapper. "
            "Permissions decide the index design, the cache key and how we call each module, so "
            "the permission model has to be agreed before any code is written.\n\n"
            "Immediate action: rotate every credential in backend/.env."
        ),
    },
    {
        "hours": 8,
        "title": "Finalise Global Search architecture for SPACE and confirm module readiness",
        "description": (
            "Date: 24-Aug-2026 (Mon).\n\n"
            "Compared the possible designs, picked one, and checked which SPACE modules are "
            "actually ready to be searched.\n\n"
            "Options compared: search from the browser, search through the module APIs, search "
            "only through Azure AI Search, or a mix of API and index.\n\n"
            "Decision - use the mix:\n"
            "- Normal records (projects, incidents, meetings) are fetched live from the existing "
            "module APIs using the user's own login, so each module applies its own rules and our "
            "search can never be more permissive than the module itself.\n"
            "- Long text (contracts, documents, drawings) goes through Azure AI Search, because no "
            "module API will full-text search a 2000 character contract clause or a PDF. Here the "
            "server always adds the permission filter, and the user cannot change it.\n"
            "- A separate Global Search service, not a module inside an existing API, because it "
            "needs its own timeout budget and its own indexing job.\n"
            "- No user search will ever touch a database directly.\n\n"
            "Why not query the SPACE databases directly (the question we will be asked):\n"
            "- One shared database login is used for everyone, so permissions cannot be applied.\n"
            "- It copies each module's business rules into search, and they drift over time.\n"
            "- It assumes the same ProjectId means the same project in every database, which is "
            "not verified - BMS already works differently.\n\n"
            "Module readiness:\n"
            "- CMS is the strongest for real search value (contract name plus eleven long text "
            "fields). Action Meeting, DMS and Project are next. Suggested starting with Action "
            "Meeting and CMS, one for each half of the design.\n"
            "- Checklist / Work Permit is not ready - no title field and the content sits in child "
            "tables. Needs the module team to expose it properly first.\n"
            "- Most result links today are empty or point to pages that do not exist, and record "
            "IDs are not stable, so both have to be fixed for clicking a result to work.\n"
            "- No tenant column exists in any of the 17 databases, so multi-tenancy needs "
            "confirmation from the backend team.\n\n"
            "Blocked on the backend team: login token details, the permission model, at least one "
            "module search API contract, and which language the new service must be written in.\n\n"
            "Can start without them: extracting the query understanding code as a clean library, "
            "and building a set of real test queries with expected results to measure relevance."
        ),
    },
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Jira Stories"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, t in enumerate(TICKETS, 2):
        ws.cell(i, 1, "").border = thin
        type_cell = ws.cell(i, 2, "Story")
        type_cell.border = thin
        type_cell.alignment = center
        comp = ws.cell(i, 3, "AI/ML")
        comp.border = thin
        comp.alignment = wrap
        title = ws.cell(i, 4, t["title"])
        title.border = thin
        title.alignment = wrap
        hours = ws.cell(i, 5, t["hours"])
        hours.border = thin
        hours.alignment = center
        desc = ws.cell(i, 6, t["description"])
        desc.border = thin
        desc.alignment = wrap
        ws.row_dimensions[i].height = 24 * (t["description"].count("\n") + 6)

    widths = [10, 12, 14, 58, 12, 95]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{len(TICKETS) + 1}"

    ws2 = wb.create_sheet("Hours by day")
    ws2["A1"] = "Worklog 21-24 Aug 2026 (Sat/Sun excluded) - 8 hours per weekday"
    ws2["A1"].font = Font(bold=True, size=13, color="1F4E79")
    ws2.merge_cells("A1:C1")
    for c, h in enumerate(["Date", "Day", "Hours"], 1):
        cell = ws2.cell(3, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
    days = [
        ("21-Aug-2026", "Fri", 8),
        ("24-Aug-2026", "Mon", 8),
    ]
    for r, row in enumerate(days, 4):
        for c, val in enumerate(row, 1):
            cell = ws2.cell(r, c, val)
            cell.border = thin
            cell.alignment = center if c > 1 else wrap

    ws2["A7"] = (
        f"Stories: {len(TICKETS)}  |  Total hours: {sum(t['hours'] for t in TICKETS)}"
        "  |  Component: AI/ML  |  Type: Story  |  Key: blank"
    )
    ws2["A7"].font = Font(bold=True)
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 10
    ws2.column_dimensions["C"].width = 10

    out = "jira-worklog-21-24-aug-global-search-analysis.xlsx"
    wb.save(out)
    print(f"Wrote {out}")
    print(f"Tickets: {len(TICKETS)} hours: {sum(t['hours'] for t in TICKETS)}")


if __name__ == "__main__":
    main()
